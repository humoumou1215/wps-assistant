# -*- coding: utf-8 -*-
"""
口径计算模块：从两个源文件算出「汇报口径」的标准答案。
被 build_target1.py / build_target2.py（生成示例行）与 verify.py（自校验）共用。

汇报口径（与源文件字段不同，需要语义映射）：
  · 全年投产次数 = 投产变更表中 变更类型 ∈ {投产上线, 紧急修复} 的记录数（优化变更不计）
  · 重大故障起数 = 故障事件表中 事件等级 ∈ {一级, I级} 的记录数
  · 平均可用率   = 月度运行表按月可用率以「交易笔数」加权；文本型百分比按数值处理，'-'/空 剔除
  · 同比         = (今年投产次数 - 去年投产次数) / 去年投产次数；去年无基线记为「新增」
  · 系统等级     = 一级→核心，二级→重要，三级→一般
"""
import json
from openpyxl import load_workbook
from pptx import Presentation

XLSX = "源A-2026年度系统运行与投产台账.xlsx"
PPTX = "源B-各研发团队年度工作条目汇编.pptx"

TEAM_FULL = {
    "研发一组": "系统研发一组",
    "研发二组": "系统研发二组",
    "研发三组": "系统研发三组",
    "数据组": "数据平台组",
    "运维组": "基础运维组",
    "安全组": "安全合规组",
}
LEVEL_CN = {"一级": "核心", "二级": "重要", "三级": "一般"}
PROD_TYPES = {"投产上线", "紧急修复"}
MAJOR_LEVELS = {"一级", "I级"}
LEVEL_NORM = {"一级": "一级", "I级": "一级", "二级": "二级", "II级": "二级", "三级": "三级"}


def _rows(ws, header_row):
    """读取数据行，自动展开合并单元格（非首格回填）。"""
    merged = {}
    for rng in ws.merged_cells.ranges:
        v = ws.cell(row=rng.min_row, column=rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged[(r, c)] = v
    out = []
    for r in range(header_row + 1, ws.max_row + 1):
        row = []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v is None and (r, c) in merged:
                v = merged[(r, c)]
            row.append(v)
        if all(x is None for x in row):
            continue
        out.append(row)
    return out


def _pct(v):
    """把各种写法的可用率归一为 float，异常返回 None。"""
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        if s in ("-", "", "—"):
            return None
        if s.endswith("%"):
            try:
                return float(s[:-1]) / 100.0
            except ValueError:
                return None
        return None
    try:
        f = float(v)
        return f if 0 < f <= 1 else None
    except (TypeError, ValueError):
        return None


def compute():
    wb = load_workbook(XLSX, data_only=True)

    # ---------------- 应用清单 ----------------
    systems = {}          # code -> dict
    order = []
    for row in _rows(wb["应用清单"], 3):
        idx, code, name, team, level, year = row[:6]
        if code is None:
            continue
        systems[code] = {
            "序号": idx, "代号": code, "全称": name, "团队简称": team,
            "团队全称": TEAM_FULL.get(team, team), "等级": LEVEL_CN.get(level, level),
            "等级源": level, "上线年份": year,
            "投产次数": 0, "重大故障": 0, "可用率": None, "_wsum": 0.0, "_w": 0,
        }
        order.append(code)

    # ---------------- 月度运行（加权可用率）----------------
    for row in _rows(wb["月度运行"], 3):
        mon, code, tx, succ, rt, av = row[:6]
        if code not in systems:
            continue
        p = _pct(av)
        if p is None:
            continue
        w = float(tx or 0)
        systems[code]["_wsum"] += p * w
        systems[code]["_w"] += w

    # ---------------- 投产变更 ----------------
    by_team_prod = {}     # 团队简称 -> {投产, 优化, 紧急, 涉资金, 回退, 影响合计}
    monthly_prod = {m: 0 for m in range(1, 13)}
    for row in _rows(wb["投产变更"], 3):
        d, code, cid, ctype, money, rollback, dur, owner = row[:8]
        if code not in systems:
            continue
        t = by_team_prod.setdefault(systems[code]["团队简称"],
                                    {"投产": 0, "优化": 0, "紧急": 0, "涉资金": 0, "回退": 0, "影响合计": 0})
        if ctype == "投产上线":
            t["投产"] += 1
            systems[code]["投产次数"] += 1
            monthly_prod[int(str(d)[5:7])] += 1
        elif ctype == "紧急修复":
            t["紧急"] += 1
            systems[code]["投产次数"] += 1
            monthly_prod[int(str(d)[5:7])] += 1
        elif ctype == "优化变更":
            t["优化"] += 1
        if money == "是":
            t["涉资金"] += 1
        if rollback in ("Y", "是"):
            t["回退"] += 1
        t["影响合计"] += int(dur or 0)

    # ---------------- 故障事件 ----------------
    by_level = {"一级": 0, "二级": 0, "三级": 0}
    cross = {}            # 等级 -> {根因: 次数}
    causes = []
    for row in _rows(wb["故障事件"], 3):
        iid, when, code, level, dur, cause, repeat = row[:7]
        if code not in systems:
            continue
        lv = LEVEL_NORM.get(level)
        if lv is None:
            continue
        if cause not in causes:
            causes.append(cause)
        by_level[lv] += 1
        cross.setdefault(lv, {}).setdefault(cause, 0)
        cross[lv][cause] += 1
        if level in MAJOR_LEVELS:
            systems[code]["重大故障"] += 1

    # ---------------- 去年基线 ----------------
    base = {}
    for row in _rows(wb["去年基线"], 3):
        code, p25, f25, a25 = row[:4]
        base[code] = {"投产": p25, "故障": f25, "可用率": _pct(a25)}

    # ---------------- 日志异常 ----------------
    for row in _rows(wb["日志异常"], 3):
        mon, code, cnt, sev = row[:4]
        if code in systems:
            systems[code].setdefault("日志异常", 0)
            systems[code]["日志异常"] += int(cnt or 0)

    # ---------------- 收口计算 ----------------
    detail = []
    for code in order:
        s = systems[code]
        s["可用率"] = (s["_wsum"] / s["_w"]) if s["_w"] else None
        b = base.get(code, {})
        prev = b.get("投产")
        if prev in (None, 0):
            s["同比"] = "新增"
            s["_sort"] = -1
        else:
            s["同比"] = (s["投产次数"] - prev) / prev
            s["_sort"] = prev
        detail.append(s)

    rank = sorted(detail, key=lambda x: (-x["投产次数"], x["序号"]))

    # ---------------- 源B：团队工作条目 ----------------
    prs = Presentation(PPTX)
    team_work = {}
    total_items = 0
    for slide in prs.slides:
        for shp in slide.shapes:
            if not shp.has_table:
                continue
            tb = shp.table
            head = [c.text.strip() for c in tb.rows[0].cells]
            if "工作条目" not in head or "责任团队" in head:
                continue
            # 团队名从该页标题文本框取
            tname = None
            for s2 in slide.shapes:
                if s2.has_text_frame and "年度工作条目" in s2.text_frame.text:
                    tname = s2.text_frame.text.split("·")[0].strip()
                    break
            if tname is None:
                tname = head[0]
            d = team_work.setdefault(tname, {"条目": 0, "已完成": 0, "进行中": 0, "延期": 0, "人天": 0})
            for r in list(tb.rows)[1:]:
                vals = [c.text.strip() for c in r.cells]
                if not any(vals):
                    continue
                st = vals[5]
                d["条目"] += 1
                if st in ("已完成", "完成"):
                    d["已完成"] += 1
                elif st == "进行中":
                    d["进行中"] += 1
                elif st == "延期":
                    d["延期"] += 1
                d["人天"] += int(vals[6] or 0)
                total_items += 1

    # ---------------- 汇总 KPI ----------------
    total_prod = sum(s["投产次数"] for s in detail)
    total_major = sum(s["重大故障"] for s in detail)
    tw = sum(s["_w"] for s in detail)
    avg_av = sum(s["_wsum"] for s in detail) / tw if tw else None

    out = {
        "systems": [{k: v for k, v in s.items() if not k.startswith("_")} for s in rank],
        "kpi": {
            "全年投产次数": total_prod,
            "重大故障起数": total_major,
            "平均可用率": avg_av,
            "完成工作条目数": sum(d["已完成"] for d in team_work.values()),
            "工作条目总数": total_items,
            "团队数": len(team_work),
            "系统数": len(detail),
        },
        "monthly_prod": monthly_prod,
        "by_team_prod": by_team_prod,
        "by_level": by_level,
        "cross": cross,
        "causes": causes,
        "team_work": team_work,
    }
    return out


if __name__ == "__main__":
    r = compute()
    print("== KPI ==")
    for k, v in r["kpi"].items():
        print(f"  {k}: {v if not isinstance(v, float) else round(v, 6)}")
    print("== 系统排名（按投产次数降序）==")
    for s in r["systems"]:
        av = f"{s['可用率']*100:.3f}%" if s["可用率"] else "-"
        yoy = f"{s['同比']*100:+.1f}%" if isinstance(s["同比"], float) else s["同比"]
        print(f"  {s['序号']:>2} {s['代号']} {s['全称']:<12} {s['团队全称']:<8} {s['等级']} "
              f"投产{s['投产次数']:>2} 重大{s['重大故障']:>2} {av:>9} 同比{yoy}")
    print("== 月度投产 ==", r["monthly_prod"])
    print("== 团队投产统计 ==")
    for k, v in r["by_team_prod"].items():
        print(f"  {k}: {v}")
    print("== 故障等级 ==", r["by_level"])
    print("== 等级×根因 ==")
    for lv, d in r["cross"].items():
        print(f"  {lv}: {d} 合计{sum(d.values())}")
    print("== 团队工作 ==")
    for k, v in r["team_work"].items():
        print(f"  {k}: {v}")
