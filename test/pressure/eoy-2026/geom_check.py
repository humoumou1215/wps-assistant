# -*- coding: utf-8 -*-
"""几何检查：PPT 形状是否溢出画布、Word 表格是否超出正文宽度。"""
import sys
from pptx import Presentation
from pptx.util import Emu
from docx import Document
from docx.shared import Cm

bad = 0
for f in ("源B-各研发团队年度工作条目汇编.pptx", "目标1-2026年度科技工作汇报.pptx"):
    prs = Presentation(f)
    W, H = prs.slide_width, prs.slide_height
    print(f"\n{f}  {W/914400:.3f}x{H/914400:.3f} in  ({len(prs.slides._sldIdLst)} 页)")
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.left is None:
                continue
            r, b = sh.left + (sh.width or 0), sh.top + (sh.height or 0)
            # 表格实际高度按行高累计
            if sh.has_table:
                th = sum((row.height or Emu(0)) for row in sh.table.rows)
                b = sh.top + max(th, sh.height or 0)
            if sh.left < -9144 or sh.top < -9144 or r > W + 9144 or b > H + 9144:
                bad += 1
                print(f"  [溢出] p{i} {sh.shape_type} right={r/914400:.2f} bottom={b/914400:.2f} "
                      f"({r/914400:.2f}>{W/914400:.2f} or {b/914400:.2f}>{H/914400:.2f})")
            elif sh.has_table:
                print(f"  p{i} 表格 {len(sh.table.rows)}行 底边 {b/914400:.2f} in")

doc = Document("目标2-2026年度生产运行分析报告.docx")
from docx.oxml.ns import qn
sec = doc.sections[0]
avail = sec.page_width - sec.left_margin - sec.right_margin
print(f"\n目标2-2026年度生产运行分析报告.docx  正文可用宽 {avail/Cm(1):.2f} cm")
for ti, t in enumerate(doc.tables, 1):
    grid = t._tbl.find(qn("w:tblGrid"))
    w = sum(int(gc.get(qn("w:w"))) for gc in grid.findall(qn("w:gridCol"))) * 635 if grid is not None else 0
    over = w - avail
    flag = f"  <-- 超宽 {over/Cm(1):+.2f} cm" if over > 9144 else "  ok"
    if over > 9144:
        bad += 1
    print(f"  表{ti}: {len(t.rows)}行 × {len(t.columns)}列  宽 {w/Cm(1):.2f} cm{flag}")

print("\n结果:", "无溢出问题 ✅" if bad == 0 else f"发现 {bad} 处问题 ❌")

sys.exit(0 if bad == 0 else 1)
