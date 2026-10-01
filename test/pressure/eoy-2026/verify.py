# -*- coding: utf-8 -*-
"""回读自检：确认 4 个文件结构完整、目标文件留白位置正确、预填样例与口径答案一致。"""
import sys, os, json
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import truth
from openpyxl import load_workbook
from docx import Document
from pptx import Presentation

R = truth.compute()
ok = True


def check(cond, msg):
    global ok
    print(("  [OK]   " if cond else "  [FAIL] ") + msg)
    if not cond:
        ok = False


print("=" * 72)
print("源A 台账")
wb = load_workbook("源A-2026年度系统运行与投产台账.xlsx")
print("  sheets:", wb.sheetnames)
for name, expect in (("应用清单", 16), ("月度运行", 192), ("投产变更", 240),
                     ("故障事件", 60), ("去年基线", 16), ("日志异常", 192)):
    ws = wb[name]
    n = sum(1 for r in range(4, ws.max_row + 1) if ws.cell(row=r, column=1).value is not None)
    check(n == expect, f"{name}: 数据行 {n}（期望 {expect}）")
ws = wb["应用清单"]
check(len(ws.merged_cells.ranges) >= 5, f"应用清单 合并单元格 {len(ws.merged_cells.ranges)} 处（责任团队列）")
ws = wb["月度运行"]
types = {}
for r in range(4, ws.max_row + 1):
    types[type(ws.cell(row=r, column=6).value).__name__] = types.get(type(ws.cell(row=r, column=6).value).__name__, 0) + 1
print("  月可用率类型分布:", types)
check(types.get("str", 0) >= 8, f"文本型/占位符可用率 {types.get('str', 0)} 处（含 \"99.97%\" 与 \"-\"）")
check(types.get("NoneType", 0) >= 1, f"空值可用率 {types.get('NoneType', 0)} 处")
ws = wb["故障事件"]
lv = {}
for r in range(4, 64):
    v = ws.cell(row=r, column=4).value
    lv[v] = lv.get(v, 0) + 1
print("  故障等级原始写法:", lv)
check("I级" in lv, "存在 I级 手写混用")
ws = wb["去年基线"]
blank = sum(1 for r in range(4, 20) if ws.cell(row=r, column=2).value is None)
check(blank == 6, f"去年基线缺基线系统 {blank} 个（期望 6）")

print("=" * 72)
print("源B 工作条目汇编")
prs = Presentation("源B-各研发团队年度工作条目汇编.pptx")
print("  slides:", len(prs.slides._sldIdLst))
tables = sum(1 for s in prs.slides for sh in s.shapes if sh.has_table)
check(tables == 7, f"PPT 内表格数 {tables}（6 个团队页 + 1 个汇总页）")
check(sum(d["条目"] for d in R["team_work"].values()) == 41, "工作条目合计 41 条")
sts = set()
for s in prs.slides:
    for sh in s.shapes:
        if not sh.has_table:
            continue
        head = [c.text.strip() for c in sh.table.rows[0].cells]
        if "工作条目" in head:
            for r in list(sh.table.rows)[1:]:
                sts.add(r.cells[5].text.strip())
check(len(sts) >= 3, f"状态字段写法: {sorted(sts)}（存在混写）")

print("=" * 72)
print("目标1 汇报 PPT")
prs = Presentation("目标1-2026年度科技工作汇报.pptx")
print("  slides:", len(prs.slides._sldIdLst))
check(len(prs.slides._sldIdLst) == 8, "共 8 页")
s5 = list(prs.slides)[4]
t5 = None
for sh in s5.shapes:
    if sh.has_table:
        t5 = sh.table
check(t5 is not None, "第 5 页存在重点系统清单表")
if t5:
    nrow, ncol = len(t5.rows), len(t5.columns)
    check(nrow == 13 and ncol == 8, f"第 5 页表格 {nrow} 行 × {ncol} 列（1 表头 + 12 容量）")
    filled = sum(1 for i in range(1, nrow) if t5.cell(i, 1).text.strip() not in ("", "【待填】"))
    check(filled == 3, f"已预填样例行 {filled} 行（应为格式样板 3 行）")
    first = [t5.cell(1, j).text for j in range(8)]
    exp = R["systems"][0]
    check(first[1] == exp["全称"] and str(first[4]) == str(exp["投产次数"]),
          f"样例首行与口径一致: {first[1]} / 投产 {first[4]}")
    print("  预填样例:", " | ".join(first))
    print(f"  容量 12 行 vs 实际需 {len(R['systems'])} 行 → 必须续页 {len(R['systems']) - 12} 行")
charts = sum(1 for s in prs.slides for sh in s.shapes if sh.has_chart)
check(charts == 1, f"图表数量 {charts}（第 4 页柱状图）")

print("=" * 72)
print("目标2 分析报告 docx")
doc = Document("目标2-2026年度生产运行分析报告.docx")
heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
print("  章节:", heads)
check(len(heads) == 5, "共 5 个一级章节")
check(len(doc.tables) == 3, f"表格数 {len(doc.tables)}")
t2, t3, t4 = doc.tables
check(len(t2.rows) == 8 and len(t2.columns) == 6, f"表2（多级表头）{len(t2.rows)} 行 × {len(t2.columns)} 列")
check(t2.cell(0, 1).text == "投产变更情况", "表2 第 1 行存在横向合并表头")
check(t2.cell(0, 0).text == "责任团队", "表2 第 1 列存在纵向合并表头")
check(len(t3.rows) == 17 and len(t3.columns) == 7, f"表3（跨页长表）{len(t3.rows)} 行 × {len(t3.columns)} 列")
check(len(t4.rows) == 7 and len(t4.columns) == 7, f"表4（团队表）{len(t4.rows)} 行 × {len(t4.columns)} 列")
todo_cnt = sum(1 for p in doc.paragraphs for r in p.runs if "【待填】" in r.text)
check(todo_cnt >= 8, f"正文待填锚点 {todo_cnt} 处")
hdr = doc.sections[0].header.paragraphs[0].text
check("文件编号" in hdr, f"页眉含文件编号占位: {hdr.strip()[:40]}…")
ftr_xml = doc.sections[0].footer.paragraphs[0]._p.xml
check("PAGE" in ftr_xml and "NUMPAGES" in ftr_xml, "页脚含 PAGE / NUMPAGES 域")
tblhdr = t3.rows[0]._tr.xml
check("tblHeader" not in tblhdr, "表3 标题行【未】设置跨页重复 → 正是待考察点")

print("=" * 72)
print("口径答案（供判定）")
k = R["kpi"]
print(f"  全年投产总次数 = {k['全年投产次数']} 次")
print(f"  重大故障起数   = {k['重大故障起数']} 起")
print(f"  系统平均可用率 = {k['平均可用率']*100:.4f}%")
print(f"  完成工作条目数 = {k['完成工作条目数']} 条（总 {k['工作条目总数']} 条）")
print("  重点系统清单（降序前 6）:")
for s in R["systems"][:6]:
    yoy = "新增" if not isinstance(s["同比"], float) else f"{s['同比']*100:+.1f}%"
    print(f"    {s['序号']:>2} {s['全称']:<12} {s['等级']} {s['团队全称']:<8} 投产{s['投产次数']:>2} "
          f"重大{s['重大故障']} {s['可用率']*100:.3f}% {yoy}")

print("=" * 72)
print("总体:", "全部通过 ✅" if ok else "存在失败项 ❌")

sys.exit(0 if ok else 1)
